"""Convert the Markdown deliverables into the .docx files the guidelines name.

    python docs/make_docx.py

Markdown stays the source of truth — it diffs, it reviews, and it is what gets edited.
The .docx files are generated artefacts for submission, regenerated whenever the Markdown
changes so the two cannot drift apart.

Supported Markdown: headings, paragraphs, bullet and numbered lists, tables, fenced code
blocks, horizontal rules, and inline bold / italic / code / links. That is the subset the
deliverables actually use; anything else passes through as plain text.
"""

from __future__ import annotations

import re
from pathlib import Path
import argparse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent

# (markdown source, output .docx, landscape?, images appended at the end)
DELIVERABLES: list[tuple[str, str, bool, list[str]]] = [
    (
        "docs/Project_Report.md",
        "docs/Project_Report.docx",
        False,
        ["architecture/AWS_Architecture.png", "architecture/System_Architecture.png"],
    ),
    ("docs/Literature_Survey.md", "docs/Literature_Survey.docx", True, []),
    ("docs/Research_Gap_Udit.md", "docs/Research_Gap_Udit.docx", False, []),
    ("docs/Objectives.md", "docs/Objectives.docx", False, []),
    ("docs/Novelty.md", "docs/Novelty.docx", False, []),
    ("dataset/dataset_description.md", "dataset/dataset_description.docx", False, []),
]

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_runs(paragraph, text: str) -> None:
    """Write ``text`` into ``paragraph``, honouring inline bold / italic / code / links."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("["):
            m = LINK.fullmatch(part)
            # Render the link text only; the URL follows in parentheses when it is external.
            if m:
                paragraph.add_run(m.group(1))
                if m.group(2).startswith("http"):
                    run = paragraph.add_run(f" ({m.group(2)})")
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
            else:
                paragraph.add_run(part)
        else:
            paragraph.add_run(part)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[str]) -> None:
    header = split_row(rows[0])
    body = [split_row(r) for r in rows[2:]]  # rows[1] is the |---|---| separator
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.autofit = True
    for cell, label in zip(table.rows[0].cells, header, strict=False):
        cell.text = ""
        add_runs(cell.paragraphs[0], label)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for row in body:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=False):
            cell.text = ""
            add_runs(cell.paragraphs[0], value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)


def convert(md_path: Path, out_path: Path, landscape: bool, images: list[str]) -> None:
    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for margin in ("left_margin", "right_margin"):
        setattr(section, margin, Inches(0.8))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            para = doc.add_paragraph()
            run = para.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue

        if (
            stripped.startswith("|")
            and i + 1 < len(lines)
            and set(lines[i + 1].strip()) <= set("|-: ")
        ):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            doc.add_paragraph()
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(level, 4))
        elif set(stripped) <= {"-", "*", "_"} and len(stripped) >= 3:
            doc.add_paragraph()  # horizontal rule
        elif re.match(r"^[-*] ", stripped):
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", stripped))
        elif stripped.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.3)
            add_runs(para, stripped.lstrip("> "))
        else:
            # Join wrapped lines of the same paragraph before writing it out.
            block = [stripped]
            i += 1
            while (
                i < len(lines)
                and lines[i].strip()
                and not re.match(r"^(#|\||```|[-*] |\d+\. |>)", lines[i].strip())
            ):
                block.append(lines[i].strip())
                i += 1
            add_runs(doc.add_paragraph(), " ".join(block))
            continue
        i += 1

    for rel in images:
        path = ROOT / rel
        if not path.exists():
            continue
        doc.add_page_break()
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(Path(rel).stem.replace("_", " "))
        run.bold = True
        width = Inches(9.0 if landscape else 6.6)
        doc.add_picture(str(path), width=width)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"{md_path.relative_to(ROOT)} -> {out_path.relative_to(ROOT)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Convert Markdown files to .docx. By default converts files listed in DELIVERABLES."
    )
    parser.add_argument("files", nargs="*", help="Markdown files to convert (workspace-relative or relative to repo root)")
    parser.add_argument("--dir", "-d", help="Directory to scan for .md files (relative to repo root)")
    parser.add_argument("--all", "-a", action="store_true", help="Convert all .md files under the repository root")
    parser.add_argument("--out", "-o", help="Output directory for generated .docx files (relative to repo root). Defaults to same directory as source files.")
    parser.add_argument("--landscape", action="store_true", help="Generate output pages in landscape orientation (applies to all outputs)")
    args = parser.parse_args()

    md_paths: list[Path] = []

    if args.all:
        md_paths = list(ROOT.rglob("*.md"))
    elif args.dir:
        md_root = (ROOT / args.dir).resolve()
        if md_root.is_dir():
            md_paths = list(md_root.rglob("*.md"))
    elif args.files:
        for f in args.files:
            p = (ROOT / f).resolve()
            if p.exists():
                md_paths.append(p)

    # If no CLI inputs were provided, fall back to DELIVERABLES list
    if not md_paths:
        for md, out, landscape, images in DELIVERABLES:
            convert(ROOT / md, ROOT / out, landscape, images)
    else:
        for md in md_paths:
            if not md.exists() or not md.suffix.lower() == ".md":
                continue
            out_dir = (ROOT / args.out) if args.out else md.parent
            out_dir.mkdir(parents=True, exist_ok=True)
            out_path = out_dir / (md.stem + ".docx")
            convert(md, out_path, args.landscape, [])


"""
Quick usage examples (run from repository root):

Convert all markdown files under the project:
    python docs/make_docx.py --all

Convert every .md in a specific folder:
    python docs/make_docx.py --dir docs

Convert specific files:
    python docs/make_docx.py docs/One.md docs/Two.md

Convert and write outputs to a single folder:
    python docs/make_docx.py --dir docs --out docs/docx_outputs

Notes:

The new CLI falls back to the original DELIVERABLES behavior when no CLI args are given.
Landscape output is controlled with --landscape (applies to all outputs produced in that run).
"""